import base64
import io
import json
import re

import openpyxl
from docx import Document
from anthropic import Anthropic

client = Anthropic()
MODEL = "claude-sonnet-4-6"

PROMPT = """You are a grocery-receipt parser. Extract every purchased food/household item from the receipt data below.

Ignore: store name, totals, tax, subtotals, tender, change, cashier, loyalty points, coupons, non-food fees.

For each item return:
- canonical_name: lowercase common name ("blueberries", not "BLUBRY PNT 6OZ")
- category: one of [produce, dairy, meat, seafood, bakery, pantry, frozen, beverage, snack, household, other]
- quantity: number (default 1 if unclear)
- unit: "oz" | "lb" | "g" | "kg" | "ct" | "pack" | "bottle" | "can" | "each"
- shelf_life_days: realistic days until spoilage
- price: number if visible, else null

Return ONLY a JSON array. No prose. No markdown fences."""

TEXT_TYPES = {
    "text/plain",
    "text/csv",
    "application/vnd.ms-excel",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/msword",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def _strip_fences(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text)
    text = re.sub(r"\s*```$", "", text)
    return text.strip()


def _extract_text(file_bytes: bytes, media_type: str) -> str:
    if media_type in ("text/plain", "text/csv"):
        return file_bytes.decode("utf-8", errors="replace")

    if media_type in ("application/vnd.ms-excel",
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"):
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                line = " | ".join(str(c) for c in row if c is not None)
                if line:
                    lines.append(line)
        return "\n".join(lines)

    if media_type in ("application/msword",
                      "application/vnd.openxmlformats-officedocument.wordprocessingml.document"):
        doc = Document(io.BytesIO(file_bytes))
        return "\n".join(p.text for p in doc.paragraphs if p.text)

    return ""


def _call_claude(messages: list) -> list[dict]:
    resp = client.messages.create(
        model=MODEL,
        max_tokens=2000,
        messages=messages,
    )
    raw = resp.content[0].text
    cleaned = _strip_fences(raw)
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        return []


def parse_receipt(file_bytes: bytes, media_type: str = "image/jpeg") -> list[dict]:
    # Text-based files: extract content and send as a text prompt
    if media_type in TEXT_TYPES:
        text = _extract_text(file_bytes, media_type)
        if not text.strip():
            return []
        return _call_claude([{
            "role": "user",
            "content": f"{PROMPT}\n\nReceipt data:\n{text}",
        }])

    # PDF: send as a document block
    if media_type == "application/pdf":
        file_data = base64.standard_b64encode(file_bytes).decode("utf-8")
        return _call_claude([{
            "role": "user",
            "content": [
                {"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": file_data}},
                {"type": "text", "text": PROMPT},
            ],
        }])

    # Images: send as an image block
    file_data = base64.standard_b64encode(file_bytes).decode("utf-8")
    return _call_claude([{
        "role": "user",
        "content": [
            {"type": "image", "source": {"type": "base64", "media_type": media_type, "data": file_data}},
            {"type": "text", "text": PROMPT},
        ],
    }])
